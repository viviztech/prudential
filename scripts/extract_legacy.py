from __future__ import annotations

from collections import Counter
from datetime import date, datetime
import json
from pathlib import Path
import re
import uuid

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "CERTIFICATE DETAILS PAS_pwd-removed.xlsx"
DATA_DIR = ROOT / "data"
OUTPUT = DATA_DIR / "legacy-certificates.json"
REPORT = DATA_DIR / "legacy-import-report.json"
NAMESPACE = uuid.UUID("62caaf7f-6aa1-4ae4-bc88-5b5bbdb45729")


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_key(value) -> str:
    return re.sub(r"[^A-Z0-9]+", " ", clean(value).upper()).strip()


def parse_date(value):
    if value in (None, "", "-"):
        return None, False
    if isinstance(value, datetime):
        return value.date().isoformat(), False
    if isinstance(value, date):
        return value.isoformat(), False

    raw = clean(value)
    formats = (
        "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y",
        "%d/%m/%y", "%d-%m-%y", "%Y/%m/%d",
    )
    for fmt in formats:
        try:
            return datetime.strptime(raw, fmt).date().isoformat(), False
        except ValueError:
            pass
    return None, True


def canonical_standard(raw_value):
    raw = normalize_key(raw_value)
    if raw.startswith("9001"):
        return "9001", "ISO 9001", "PASQM"
    if raw.startswith("14001"):
        return "14001", "ISO 14001", "PASEM"
    if raw.startswith("45001"):
        return "45001", "ISO 45001", "PASOH"
    if raw.startswith("18001"):
        return "18001", "OHSAS 18001", "PASOH"
    if raw.startswith("22000"):
        return "22000", "ISO 22000", "PASFS"
    if raw.startswith("27001") or raw.startswith("27000"):
        return "27001", "ISO 27000", "PASIS"
    if raw.startswith("13485"):
        return "13485", "ISO 13485", "PASMD"
    if raw.startswith("20001") or raw.startswith("20000"):
        return "20001", "ISO 20001", "PASIT"
    if raw.startswith("17024"):
        return "17024", "ISO 17024:2017", "PASCS"
    if raw.startswith("SA 8000") or raw.startswith("SA8000"):
        return "SA8000", "SA 8000", "PASSA"
    if raw == "GMP":
        return "GMP", "GMP", "PASGM"
    if raw == "HACCP":
        return "HACCP", "HACCP", "PASHA"
    if raw == "CE":
        return "CE", "CE", "PASCE"
    if raw == "ROHS":
        return "ROHS", "ROHS", "PASRO"
    if raw == "GREEN":
        return "GREEN", "GREEN", "PASGR"
    return "OTHER", "Other / Unspecified", "PASOT"


def yes(value) -> bool:
    return clean(value).upper() == "YES"


def main():
    workbook = load_workbook(SOURCE, data_only=True, read_only=False)
    sheet = workbook["Clients"]
    populated_rows = [
        row for row in range(2, sheet.max_row + 1)
        if clean(sheet.cell(row, 4).value)
    ]

    raw_numbers = [clean(sheet.cell(row, 15).value).upper() for row in populated_rows]
    usable_numbers = [number for number in raw_numbers if number and number != "-"]
    number_counts = Counter(usable_numbers)
    duplicate_numbers = {number for number, count in number_counts.items() if count > 1}

    records = []
    issues = []
    standards = Counter()
    company_ids = set()
    date_columns = {
        "application_date": 3,
        "draft_date": 10,
        "issue_date": 11,
        "first_surveillance_date": 12,
        "second_surveillance_date": 13,
        "expiry_date": 14,
    }

    for row in populated_rows:
        company_name = clean(sheet.cell(row, 4).value)
        address = clean(sheet.cell(row, 5).value)
        company_key = f"{normalize_key(company_name)}|{normalize_key(address)}"
        company_id = str(uuid.uuid5(NAMESPACE, f"company:{company_key}"))
        company_ids.add(company_id)
        certificate_id = str(uuid.uuid5(NAMESPACE, f"certificate:clients-row-{row}"))

        parsed_dates = {}
        for field, column in date_columns.items():
            parsed, failed = parse_date(sheet.cell(row, column).value)
            parsed_dates[field] = parsed
            if failed:
                issues.append({
                    "row": row,
                    "type": "unparsed_date",
                    "field": field,
                    "value": clean(sheet.cell(row, column).value),
                })

        original_standard = clean(sheet.cell(row, 9).value)
        standard_code, standard_name, certificate_prefix = canonical_standard(original_standard)
        standards[standard_code] += 1
        raw_number = clean(sheet.cell(row, 15).value).upper()
        public_number = raw_number if raw_number and raw_number != "-" and raw_number not in duplicate_numbers else None
        hold = clean(sheet.cell(row, 20).value).upper() == "HOLD"
        printed = yes(sheet.cell(row, 17).value)

        if hold:
            status = "on_hold"
        elif public_number and printed:
            status = "printed"
        elif public_number:
            status = "issued"
        elif clean(sheet.cell(row, 16).value).upper() == "COMPLETED":
            status = "approved"
        else:
            status = "draft_created"

        if raw_number in duplicate_numbers:
            issues.append({"row": row, "type": "duplicate_certificate_number", "value": raw_number})
        elif not raw_number or raw_number == "-":
            issues.append({"row": row, "type": "missing_certificate_number", "value": raw_number or None})
        if standard_code == "OTHER":
            issues.append({"row": row, "type": "unmapped_standard", "value": original_standard})

        records.append({
            "source_sheet": "Clients",
            "source_row": row,
            "certificate_id": certificate_id,
            "company_id": company_id,
            "company_name": company_name,
            "address": address,
            "scope": clean(sheet.cell(row, 6).value),
            "contact_person": clean(sheet.cell(row, 7).value),
            "mobile": clean(sheet.cell(row, 8).value),
            "email": None,
            "associate_name": clean(sheet.cell(row, 2).value),
            "standard_code": standard_code,
            "standard_name": standard_name,
            "certificate_prefix": certificate_prefix,
            "original_standard": original_standard,
            "certificate_number": public_number,
            "legacy_certificate_number": raw_number or None,
            "status": status,
            "printed": printed,
            "delivered": yes(sheet.cell(row, 18).value),
            "activated": yes(sheet.cell(row, 19).value),
            **parsed_dates,
            "third_surveillance_date": parsed_dates["expiry_date"],
        })

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps({"source": SOURCE.name, "records": records}, ensure_ascii=False, indent=2), encoding="utf-8")
    report = {
        "source": SOURCE.name,
        "source_sheet": "Clients",
        "records": len(records),
        "unique_companies": len(company_ids),
        "public_certificate_numbers": sum(1 for record in records if record["certificate_number"]),
        "duplicate_certificate_numbers": sorted(duplicate_numbers),
        "records_with_duplicate_numbers": sum(1 for record in records if record["legacy_certificate_number"] in duplicate_numbers),
        "records_missing_certificate_number": sum(1 for record in records if not record["legacy_certificate_number"] or record["legacy_certificate_number"] == "-"),
        "standards": dict(sorted(standards.items())),
        "issue_count": len(issues),
        "issues": issues,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: report[key] for key in ("records", "unique_companies", "public_certificate_numbers", "records_with_duplicate_numbers", "records_missing_certificate_number", "issue_count")}, indent=2))


if __name__ == "__main__":
    main()
